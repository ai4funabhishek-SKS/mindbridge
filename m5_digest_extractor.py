"""
M5 Phase 2 — Digest extractor for the 257 unprocessed files.
Produces Processed_Assets/m5_digests.json: one compact digest per file so
classification batches can be accurate without loading full binaries.
Read-only over Raw_copies/.
"""
import json
import re
from pathlib import Path

ROOT = Path(__file__).parent.resolve()
QUEUE = ROOT / "Processed_Assets" / "m5_remaining_queue.json"
OUT = ROOT / "Processed_Assets" / "m5_digests.json"

CAP = 1400


def clamp(t, n=CAP):
    t = re.sub(r"[ \t]+", " ", t)
    return t.strip()[:n]


def digest_text(p):
    raw = p.read_text(encoding="utf-8", errors="ignore")
    headings = re.findall(r"^#{1,3}\s+(.{3,80})$", raw, re.M)[:8]
    code_blocks = len(re.findall(r"```", raw)) // 2
    urls = len(set(re.findall(r"https?://\S{8,}", raw)))
    promptish = len(re.findall(r"(?i)(?:you are|act as|your (?:task|role)|system prompt|<\w+>)", raw))
    return {
        "kind": "text",
        "headings": [h.strip() for h in headings],
        "counts": {"code_blocks": code_blocks, "urls": urls, "prompt_markers": promptish},
        "preview": clamp(raw),
    }


def digest_json(p):
    try:
        d = json.loads(p.read_text(encoding="utf-8", errors="ignore"))
    except Exception as e:
        return {"kind": "json_invalid", "preview": clamp(str(e))}
    if isinstance(d, dict) and "nodes" in d and isinstance(d["nodes"], list):
        nodes = d["nodes"]
        types = sorted({n.get("type", "?").split(".")[-1] for n in nodes})[:15]
        names = [n.get("name", "") for n in nodes][:12]
        return {
            "kind": "n8n_workflow",
            "node_count": len(nodes),
            "node_types": types,
            "node_names": names,
            "has_credentials_flag": bool(re.search(
                r"(?i)(credential|apikey|api_key|password|token)",
                json.dumps(d)[:20000])),
        }
    if isinstance(d, list):
        sample = json.dumps(d[:2], ensure_ascii=False)
        return {"kind": "json_list", "len": len(d), "preview": clamp(sample)}
    keys = list(d.keys())[:20] if isinstance(d, dict) else ["<scalar>"]
    return {"kind": "json_dict", "keys": keys, "preview": clamp(json.dumps(d, ensure_ascii=False))}


def digest_pdf(p):
    try:
        import fitz
        doc = fitz.open(p)
        npages = doc.page_count
        txt = ""
        for i in range(min(3, npages)):
            txt += doc[i].get_text()
        doc.close()
        return {"kind": "pdf", "pages": npages, "preview": clamp(txt)}
    except Exception as e:
        return {"kind": "pdf_error", "error": str(e)[:120]}


def digest_docx(p):
    try:
        import docx
        d = docx.Document(p)
        paras = [x.text.strip() for x in d.paragraphs if x.text.strip()][:12]
        return {"kind": "docx", "paragraphs": len(d.paragraphs),
                "tables": len(d.tables), "preview": clamp(" | ".join(paras))}
    except Exception as e:
        return {"kind": "docx_error", "error": str(e)[:120]}


def digest_xlsx(p):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True)
        sheets = {s: wb[s].max_row for s in wb.sheetnames[:10]}
        wb.close()
        return {"kind": "xlsx", "sheets": sheets}
    except Exception as e:
        return {"kind": "xlsx_error", "error": str(e)[:120]}


DIGESTERS = {
    ".md": digest_text, ".txt": digest_text,
    ".json": digest_json,
    ".pdf": digest_pdf, ".docx": digest_docx, ".xlsx": digest_xlsx,
}

def main():
    queue = json.loads(QUEUE.read_text(encoding="utf-8"))
    out = {}
    errors = []
    for rel in queue:
        p = ROOT / rel
        ext = p.suffix.lower()
        fn = DIGESTERS.get(ext)
        if fn is None:
            out[rel] = {"kind": f"unsupported{ext}", "size": p.stat().st_size}
            continue
        try:
            out[rel] = fn(p)
        except Exception as e:
            errors.append(rel)
            out[rel] = {"kind": "error", "error": str(e)[:160]}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    from collections import Counter
    print("digests:", len(out), "| hard errors:", len(errors))
    print(dict(Counter(v["kind"] for v in out.values())))

if __name__ == "__main__":
    main()
