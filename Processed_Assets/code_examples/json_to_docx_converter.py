import json
import os
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

def convert_json_to_docx(json_file_path, docx_file_path=None):
    """Convert LLM conversation JSON to DOCX format"""

    # Load JSON data
    with open(json_file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    conversations = data.get('conversations', [])
    summary = data.get('summary', {})

    # Create DOCX file path if not provided
    if docx_file_path is None:
        json_path = Path(json_file_path)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        docx_file_path = json_path.parent / f"{json_path.stem}_{timestamp}.docx"

    # Create document
    doc = Document()

    # Add title
    title = doc.add_heading(f'{summary.get("llm_platform", "LLM")} Conversation Export', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Add summary information
    doc.add_heading('Export Summary', level=1)
    summary_para = doc.add_paragraph()
    summary_para.add_run(f'Platform: {summary.get("llm_platform", "Unknown")}\n')
    summary_para.add_run(f'Total Conversations: {summary.get("total_conversations", 0)}\n')
    summary_para.add_run(f'Export Date: {summary.get("extraction_end", datetime.now().isoformat())}\n')
    summary_para.add_run(f'Duration: {summary.get("extraction_duration_seconds", 0):.2f} seconds\n')

    # Add table of contents
    doc.add_heading('Table of Contents', level=1)
    toc_para = doc.add_paragraph()
    for i, conv in enumerate(conversations, 1):
        toc_para.add_run(f'{i}. {conv.get("title", f"Conversation {i}")}\n')

    # Add each conversation
    for i, conv in enumerate(conversations, 1):
        doc.add_page_break()

        # Conversation header
        conv_title = doc.add_heading(f'Conversation {i}: {conv.get("title", "Untitled")}', level=1)

        # Conversation metadata
        meta_para = doc.add_paragraph()
        meta_para.add_run(f'ID: {conv.get("id", "N/A")}\n')
        meta_para.add_run(f'Source: {conv.get("source", "Unknown")}\n')
        meta_para.add_run(f'Export Date: {conv.get("export_date", datetime.now().isoformat())}\n')
        if conv.get('creation_date'):
            meta_para.add_run(f'Created: {conv.get("creation_date")}\n')
        if conv.get('url'):
            meta_para.add_run(f'URL: {conv.get("url")}\n')

        # Add conversation thread
        thread = conv.get('thread', [])
        if thread:
            doc.add_heading('Conversation Thread', level=2)

            for j, message in enumerate(thread, 1):
                # Message header
                role = message.get('role', 'unknown')
                timestamp = message.get('timestamp', '')

                msg_header = doc.add_paragraph()
                msg_header.add_run(f'{j}. {role.title()}: ').bold = True
                if timestamp:
                    msg_header.add_run(f'({timestamp})')

                # Message content
                content = message.get('content', '')
                if content:
                    content_para = doc.add_paragraph(content)
                    content_para.style = 'Normal'

                    # Add some spacing
                    doc.add_paragraph('')

        else:
            doc.add_paragraph('No messages found in this conversation.')

    # Save document
    doc.save(docx_file_path)
    print(f'✓ Successfully converted {json_file_path} to {docx_file_path}')
    return docx_file_path

def batch_convert_json_to_docx(source_dir=None, target_dir=None):
    """Convert all JSON conversation files in a directory to DOCX"""

    if source_dir is None:
        source_dir = Path('c:/ai-stack/knowledge Base/my research/sources/raw_data')

    if target_dir is None:
        target_dir = source_dir

    source_path = Path(source_dir)
    target_path = Path(target_dir)
    target_path.mkdir(exist_ok=True)

    # Find all JSON files that contain conversations
    json_files = list(source_path.glob('*conversations*.json'))

    if not json_files:
        print(f'No conversation JSON files found in {source_path}')
        return []

    converted_files = []
    for json_file in json_files:
        try:
            docx_filename = json_file.stem + '.docx'
            docx_path = target_path / docx_filename

            convert_json_to_docx(json_file, docx_path)
            converted_files.append(docx_path)

        except Exception as e:
            print(f'Error converting {json_file}: {str(e)}')
            continue

    print(f'\n✓ Batch conversion complete! Converted {len(converted_files)} files.')
    return converted_files

def main():
    import sys

    if len(sys.argv) < 2:
        print("Usage:")
        print("  python json_to_docx_converter.py <json_file_path> [docx_file_path]")
        print("  python json_to_docx_converter.py --batch [source_dir] [target_dir]")
        return

    if sys.argv[1] == '--batch':
        source_dir = sys.argv[2] if len(sys.argv) > 2 else None
        target_dir = sys.argv[3] if len(sys.argv) > 3 else None
        batch_convert_json_to_docx(source_dir, target_dir)
    else:
        json_file = sys.argv[1]
        docx_file = sys.argv[2] if len(sys.argv) > 2 else None
        convert_json_to_docx(json_file, docx_file)

if __name__ == '__main__':
    main()
