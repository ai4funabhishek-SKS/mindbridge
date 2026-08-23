import os
import sys
import json
import re
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import fitz # PyMuPDF
import docx

BASE_DIR = Path(r"C:\KnowledgeBase_Staging")
if not BASE_DIR.exists():
    BASE_DIR = Path(__file__).resolve().parent

RAW_DIR = BASE_DIR / "Raw_copies"
TOPICS_DIR = BASE_DIR / "03_Topics_and_Projects"
OUTPUT_DIR = BASE_DIR / "00_Index_and_Dashboards"
EXCEL_OUTPUT = OUTPUT_DIR / "Master_Second_Brain_Knowledge_Mine.xlsx"
EXCEL_OUTPUT_ALT = OUTPUT_DIR / "Master_Second_Brain_Dashboard.xlsx"
HTML_OUTPUT = OUTPUT_DIR / "Interactive_Knowledge_Mine_Navigator.html"

CANONICAL_DOMAINS = [
    "01_AI_Agents_and_Autonomous_Systems",
    "02_Workflow_Automations_n8n",
    "03_Prompt_Engineering_and_Frameworks",
    "04_Mobile_and_Android_Ecosystem",
    "05_Code_Snippets_and_Dev_Tools",
    "06_Quantitative_Finance_and_Trading",
    "07_Content_Creation_and_Media_Engines",
    "08_Business_Operations_and_Subbroking",
    "09_LLM_Conversations_and_Syntheses",
    "10_R_and_D_Incubator_Half_Baked_Ideas"
]

def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'[\r\n\t]+', ' ', text)
    text = re.sub(r'\s{2,}', ' ', text)
    return text.strip()

def inspect_json(file_path: Path):
    title = file_path.stem
    topic = "AI & Automation"
    problem_solved = "Automated data workflow / configuration"
    workability = "Deployable Template"
    reusability = "High Reusability (Direct Mine)"
    tags = []
    summary = ""
    
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            data = json.load(f)
            
        if isinstance(data, dict):
            if "nodes" in data or "connections" in data:
                nodes = data.get("nodes", [])
                node_types = []
                node_names = []
                for n in nodes:
                    if isinstance(n, dict):
                        ntype = n.get("type", "").split(".")[-1]
                        nname = n.get("name", "")
                        if ntype: node_types.append(ntype)
                        if nname: node_names.append(nname)
                
                unique_types = list(dict.fromkeys(node_types))[:8]
                tags = ["n8n", "Workflow Automation"] + unique_types[:4]
                
                name_val = data.get("name", "")
                if name_val and name_val != "My workflow":
                    title = f"n8n: {name_val}"
                else:
                    title = f"n8n Workflow ({file_path.stem})"
                
                combined_info = (file_path.stem + " " + " ".join(node_names) + " " + " ".join(node_types)).lower()
                
                if any(k in combined_info for k in ["video", "youtube", "faceless", "audio", "4k"]):
                    topic = "Content Creation & Video"
                    problem_solved = "Automates video creation, YouTube publishing, or media processing"
                elif any(k in combined_info for k in ["trading", "stock", "crypto", "finance", "forex", "market"]):
                    topic = "Trading & Financial Systems"
                    problem_solved = "Automates financial market data retrieval, alerts, or trading signals"
                elif any(k in combined_info for k in ["pinecone", "rag", "vector", "embedding", "document"]):
                    topic = "AI & RAG Knowledge Systems"
                    problem_solved = "Indexes documents and vectors into Pinecone / Vector DB for RAG querying"
                elif any(k in combined_info for k in ["email", "gmail", "label"]):
                    topic = "AI & Productivity Automation"
                    problem_solved = "Automates email triage, classification, and drafting with AI"
                elif any(k in combined_info for k in ["social", "post", "linkedin", "twitter", "instagram"]):
                    topic = "Content Creation & Social Media"
                    problem_solved = "Multi-platform automated social media distribution and scheduling"
                elif any(k in combined_info for k in ["call", "vapi", "phone", "voice", "twilio"]):
                    topic = "AI Voice & Telephony"
                    problem_solved = "Automates outbound/inbound phone calls with AI voice agents"
                else:
                    topic = "AI & Workflow Automation"
                    problem_solved = f"Automates pipeline utilizing {', '.join(unique_types[:3]) if unique_types else 'n8n nodes'}"
                
                workability = "Production-Ready Workflow"
                summary = f"n8n Workflow with {len(nodes)} nodes ({', '.join(node_names[:4])}). Integrates: {', '.join(unique_types[:4])}."
                
                raw_str = json.dumps(data).lower()
                if any(k in raw_str for k in ["api_key", "secret", "bearer", "password", "token"]) and any(k in raw_str for k in ["sk-", "ghp_", "eyj"]):
                    reusability = "Sensitive (Private Vault)"
                    tags.append("Contains Keys")
                
            elif "system_prompt" in data or "tools" in data or "role" in data:
                topic = "AI Agents & Autonomous Systems"
                agent_role = data.get("role", data.get("name", file_path.stem))
                title = f"AI Agent: {agent_role}"
                problem_solved = f"Autonomous agent specialized in {agent_role}"
                workability = "Deployable Agent Config"
                tags = ["AI Agent", "System Architecture", "Autonomous Agent"]
                summary = f"Agent configuration for {agent_role}. Tools: {list(data.get('tools', {}).keys()) if isinstance(data.get('tools'), dict) else 'Custom tools'}."
            else:
                title = f"JSON Config: {file_path.stem}"
                topic = "Developer Tooling & Configs"
                problem_solved = "Structured data configuration or parameters"
                workability = "Deployable Template"
                tags = ["JSON", "Config Data"]
                summary = f"JSON structure with keys: {list(data.keys())[:6]}."
        elif isinstance(data, list):
            title = f"Dataset: {file_path.stem} ({len(data)} items)"
            topic = "Datasets & Tables"
            problem_solved = "Structured tabular/list data repository"
            workability = "Reference Data"
            tags = ["Dataset", "JSON Array"]
            summary = f"Array containing {len(data)} data records."
    except Exception as e:
        title = file_path.name
        summary = f"JSON file note: {str(e)[:100]}"
        tags = ["JSON"]
        
    return {
        "title": title,
        "topic": topic,
        "problem_solved": problem_solved,
        "workability": workability,
        "reusability": reusability,
        "tags": ", ".join(tags),
        "summary": clean_text(summary)
    }

def inspect_pdf(file_path: Path):
    title = file_path.stem
    topic = "Reference Documents"
    problem_solved = "Documentation & Reference Knowledge"
    workability = "Reference Documentation"
    reusability = "Reference Knowledge"
    tags = ["PDF"]
    summary = ""
    
    try:
        doc = fitz.open(file_path)
        num_pages = len(doc)
        text_samples = []
        for i in range(min(3, num_pages)):
            try:
                text_samples.append(doc[i].get_text())
            except Exception:
                pass
        doc.close()
        
        full_sample = " ".join(text_samples)
        sample_lower = full_sample.lower()
        fn_lower = file_path.name.lower()
        
        if any(k in fn_lower or k in sample_lower for k in ["statement", "account statement", "hdfc", "sbi", "icici", "axis", "kotak", "bank", "credit card", "debit", "balance"]):
            topic = "Personal Finance & Bank Statements"
            problem_solved = "Financial transaction record & bank statement"
            workability = "Private Reference Data"
            reusability = "Sensitive (Private Vault)"
            tags = ["Finance", "Bank Statement", "Private Vault"]
            title = f"Bank Statement: {file_path.stem}"
            summary = f"Financial statement document ({num_pages} pages). Contains transactional accounting records."
        elif any(k in fn_lower or k in sample_lower for k in ["adhar", "aadhar", "pan card", "passport", "driving license", "voter id"]):
            topic = "Personal Identity & KYC"
            problem_solved = "Official identity and KYC verification"
            workability = "Private Identification"
            reusability = "Sensitive (Private Vault)"
            tags = ["KYC", "Identity", "Private Vault"]
            title = f"Identity Document: {file_path.stem}"
            summary = "Personal identity / government ID documentation."
        elif any(k in fn_lower or k in sample_lower for k in ["prompt engineering", "whitepaper", "chatgpt", "deepseek", "llm", "transformer", "generative ai"]):
            topic = "Prompt Engineering & AI Research"
            problem_solved = "Comprehensive guide and techniques for advanced prompt engineering and LLM orchestration"
            workability = "Framework & Best Practices"
            reusability = "Must Keep - High Reusability"
            tags = ["Prompt Engineering", "AI Research", "Whitepaper", "LLM Best Practices"]
            title = f"AI Whitepaper: {file_path.stem[:45]}"
            summary = f"Comprehensive research whitepaper / guide ({num_pages} pages). Covers prompt optimization, reasoning architectures, and LLM implementation."
        else:
            summary = f"PDF Document ({num_pages} pages). Excerpt: {clean_text(full_sample[:200])}"
            tags.append("Documentation")
    except Exception as e:
        summary = f"PDF note: {str(e)[:100]}"
        
    return {
        "title": title,
        "topic": topic,
        "problem_solved": problem_solved,
        "workability": workability,
        "reusability": reusability,
        "tags": ", ".join(tags),
        "summary": clean_text(summary)
    }

def inspect_markdown_and_text(file_path: Path):
    title = file_path.stem
    topic = "Notes & Brainstorms"
    problem_solved = "Conceptual ideas and reference material"
    workability = "Conceptual Idea / R&D"
    reusability = "Half-Baked Idea (R&D Mine)"
    tags = ["Markdown" if file_path.suffix.lower() == '.md' else "Text"]
    summary = ""
    
    try:
        content = file_path.read_text(encoding='utf-8', errors='ignore')
        content_lower = content.lower()
        fn_lower = file_path.name.lower()
        lines = [l.strip() for l in content.splitlines() if l.strip()]
        
        first_heading = ""
        for line in lines[:5]:
            if line.startswith("# "):
                first_heading = line.replace("# ", "").strip()
                break
        if first_heading:
            title = first_heading[:60]
            
        if any(k in fn_lower or k in content_lower for k in ["prompt", "system prompt", "act as", "you are an expert", "role:", "instruction:"]) or "prompts" in str(file_path).lower():
            topic = "Prompt Engineering & System Prompts"
            problem_solved = "Ready-to-use system prompts and few-shot templates for AI models"
            workability = "Deployable Prompt Template"
            reusability = "Must Keep - High Reusability"
            tags.extend(["Prompts", "System Instructions", "AI Templates"])
            title = f"Prompt: {first_heading if first_heading else file_path.stem}"
            summary = f"Prompt collection/template. Key instructions: {clean_text(content[:250])}"
        elif any(k in fn_lower or k in content_lower for k in ["trading", "subbroker", "equity", "nifty", "banknifty", "options", "forex", "sma", "rsi", "indicator"]):
            topic = "Trading & Financial Systems"
            problem_solved = "Trading rules, quantitative algorithms, or subbroking business logic"
            workability = "Algorithmic Rules & Logic"
            reusability = "Must Keep - High Reusability"
            tags.extend(["Trading", "Finance", "TradingView", "Algorithms"])
            title = f"Trading Logic: {first_heading if first_heading else file_path.stem}"
            summary = f"Trading system / market research notes. Excerpt: {clean_text(content[:250])}"
        elif any(k in fn_lower for k in ["! inbox", "inbox", "idea", "concept", "draft", "todo", "temp", "rough"]):
            topic = "Ideas & Half-Baked Concepts"
            problem_solved = "Early-stage architecture draft and brainstorming mine for future venture R&D"
            workability = "Conceptual Idea / R&D"
            reusability = "Half-Baked Idea (R&D Mine)"
            tags.extend(["R&D Mine", "Idea Draft", "Brainstorm"])
            title = f"Draft Idea: {first_heading if first_heading else file_path.stem}"
            summary = f"Brainstorming notes requiring synthesis. Excerpt: {clean_text(content[:250])}"
        elif fn_lower.startswith("agent.") or any(k in content_lower for k in ["def ", "class ", "import ", "function"]):
            topic = "AI Agents & Code Snippets"
            problem_solved = "Reusable agent behavior modules, functions, and communication contracts"
            workability = "Production-Ready Snippet"
            reusability = "Must Keep - High Reusability"
            tags.extend(["Code Snippet", "Agent Behavior", "Developer Tool"])
            title = f"Agent Module: {file_path.name}"
            summary = f"Agent behavioral or code logic definition. Excerpt: {clean_text(content[:250])}"
        else:
            topic = "General Knowledge & Notes"
            problem_solved = "Knowledge base reference material"
            workability = "Reference Documentation"
            reusability = "Reference Knowledge"
            tags.append("Reference")
            summary = f"Notes excerpt: {clean_text(content[:250])}"
    except Exception as e:
        summary = f"Text parsing error: {str(e)[:100]}"
        
    return {
        "title": title,
        "topic": topic,
        "problem_solved": problem_solved,
        "workability": workability,
        "reusability": reusability,
        "tags": ", ".join(tags),
        "summary": clean_text(summary)
    }

def inspect_docx(file_path: Path):
    title = file_path.stem
    topic = "Research & Reports"
    problem_solved = "Structured documentation & synthesis"
    workability = "Reference Documentation"
    reusability = "Must Keep - High Reusability"
    tags = ["Word Document"]
    summary = ""
    
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = " ".join(paragraphs)
        fn_lower = file_path.name.lower()
        
        if "cv" in fn_lower or "resume" in fn_lower:
            topic = "Personal Identity & Career"
            problem_solved = "Professional curriculum vitae & skillset profile"
            workability = "Production-Ready Profile"
            reusability = "Must Keep - High Reusability"
            tags.extend(["Resume", "CV", "Profile"])
            title = f"CV / Resume: {file_path.stem}"
            summary = "Curriculum Vitae detailing professional experience and technical competencies."
        elif "deepseek" in fn_lower or "research" in fn_lower:
            topic = "AI Research & DeepSeek Intelligence"
            problem_solved = "Synthesis of DeepSeek LLM capabilities, architecture, and deployment strategies"
            workability = "Synthesis Report"
            reusability = "Must Keep - High Reusability"
            tags.extend(["AI Research", "DeepSeek", "Intelligence Report"])
            title = f"Research Report: {file_path.stem}"
            summary = f"Comprehensive research synthesis report ({len(paragraphs)} sections). Excerpt: {clean_text(full_text[:250])}"
        elif "n8n" in fn_lower or "workflow" in fn_lower:
            topic = "AI & Workflow Automation Documentation"
            problem_solved = "Comprehensive n8n node architecture and workflow creation guide"
            workability = "Production Documentation"
            reusability = "Must Keep - High Reusability"
            tags.extend(["n8n", "Workflow Guide", "Node Docs"])
            title = f"Workflow Documentation: {file_path.stem}"
            summary = f"n8n workflow and node reference documentation ({len(paragraphs)} sections)."
        else:
            summary = f"Word document ({len(paragraphs)} paragraphs). Excerpt: {clean_text(full_text[:250])}"
    except Exception as e:
        summary = f"DOCX parsing note: {str(e)[:100]}"
        
    return {
        "title": title,
        "topic": topic,
        "problem_solved": problem_solved,
        "workability": workability,
        "reusability": reusability,
        "tags": ", ".join(tags),
        "summary": clean_text(summary)
    }

def inspect_xlsx(file_path: Path):
    title = file_path.stem
    topic = "Financial & Tabular Data"
    problem_solved = "Tabular data, client records, or account statements"
    workability = "Data Table"
    reusability = "Reference Knowledge"
    tags = ["Spreadsheet", "Excel", "Data Table"]
    summary = ""
    
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        fn_lower = file_path.name.lower()
        
        if any(k in fn_lower for k in ["statement", "bank", "gain", "loss", "holding", "report"]):
            topic = "Personal Finance & Bank Statements"
            problem_solved = "Financial transaction ledger & account summary"
            workability = "Private Financial Data"
            reusability = "Sensitive (Private Vault)"
            tags.extend(["Bank Ledger", "Private Vault"])
            title = f"Financial Sheet: {file_path.stem}"
            summary = f"Excel financial statements with sheets: {', '.join(sheet_names[:6])}."
        else:
            title = f"Data Sheet: {file_path.stem}"
            summary = f"Excel spreadsheet with {len(sheet_names)} sheets: {', '.join(sheet_names[:6])}."
    except Exception as e:
        summary = f"Excel parsing note (may be password-protected): {str(e)[:100]}"
        if "statement" in file_path.name.lower():
            reusability = "Sensitive (Private Vault)"
            topic = "Personal Finance & Bank Statements"
            workability = "Encrypted Financial Statement"
            
    return {
        "title": title,
        "topic": topic,
        "problem_solved": problem_solved,
        "workability": workability,
        "reusability": reusability,
        "tags": ", ".join(tags),
        "summary": clean_text(summary)
    }

def process_all_files() -> pd.DataFrame:
    records = []
    print(f"Scanning raw repository at {RAW_DIR}...")
    
    # 1. Scan Raw_copies/
    for root, _, files in os.walk(RAW_DIR):
        for f in files:
            file_path = Path(root) / f
            ext = file_path.suffix.lower()
            size_kb = round(file_path.stat().st_size / 1024, 2)
            rel_path = str(file_path.relative_to(BASE_DIR)).replace('\\', '/')
            
            if ext == '.json':
                details = inspect_json(file_path)
                asset_format = "JSON / n8n Workflow" if "n8n" in details["tags"] else "JSON Config"
            elif ext == '.pdf':
                details = inspect_pdf(file_path)
                asset_format = "PDF Document"
            elif ext in ['.md', '.txt']:
                details = inspect_markdown_and_text(file_path)
                asset_format = "Markdown Note" if ext == '.md' else "Text File"
            elif ext in ['.docx', '.doc']:
                details = inspect_docx(file_path)
                asset_format = "Word Document"
            elif ext in ['.xlsx', '.xls', '.csv']:
                details = inspect_xlsx(file_path)
                asset_format = "Spreadsheet"
            else:
                asset_format = "Other Asset"
                details = {
                    "title": file_path.stem,
                    "topic": "General Asset",
                    "problem_solved": "General asset storage",
                    "workability": "Reference Data",
                    "reusability": "Reference Knowledge",
                    "tags": ext,
                    "summary": f"File of format {ext}"
                }
                
            records.append({
                "File Name": file_path.name,
                "Relative Path": rel_path,
                "Asset Title": details["title"],
                "Asset Format": asset_format,
                "Size (KB)": size_kb,
                "Primary Topic": details["topic"],
                "Problem Solved": details["problem_solved"],
                "Workability Status": details["workability"],
                "Actionable Rating": details["reusability"],
                "Tags & Technologies": details["tags"],
                "Content Summary & Logic": details["summary"]
            })

    # 2. Scan Synthesized Golden Master Playbooks in 03_Topics_and_Projects/
    if TOPICS_DIR.exists():
        for root, _, files in os.walk(TOPICS_DIR):
            for f in files:
                if not f.endswith(".md"):
                    continue
                file_path = Path(root) / f
                size_kb = round(file_path.stat().st_size / 1024, 2)
                rel_path = str(file_path.relative_to(BASE_DIR)).replace('\\', '/')
                
                content = file_path.read_text(encoding='utf-8', errors='ignore')
                domain_name = file_path.parent.name
                
                # Extract title from H1 or filename
                title = file_path.stem.replace("GM_", "").replace("_", " ")
                for line in content.splitlines()[:15]:
                    if line.startswith("# "):
                        title = line.replace("# ", "").strip()
                        break
                        
                records.append({
                    "File Name": file_path.name,
                    "Relative Path": rel_path,
                    "Asset Title": f"🏆 {title}",
                    "Asset Format": "Golden Master Playbook",
                    "Size (KB)": size_kb,
                    "Primary Topic": domain_name,
                    "Problem Solved": f"Institutional synthesized Golden Master Playbook for {domain_name}",
                    "Workability Status": "Production-Ready",
                    "Actionable Rating": "Must Keep - High Reusability",
                    "Tags & Technologies": f"Golden Master, {domain_name}, Institutional Architecture",
                    "Content Summary & Logic": clean_text(content[:300])
                })
            
    df = pd.DataFrame(records)
    print(f"Deep inspection complete! Processed {len(df)} assets.")
    return df

def style_excel_sheet(ws, is_kpi=False):
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )
    
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 22
        is_alt = (row % 2 == 0)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = border_thin
            if is_alt and not is_kpi:
                cell.fill = alt_fill
            cell.alignment = Alignment(vertical="center", wrap_text=(col in [7, 10, 11]))
            
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

def build_enhanced_excel(df: pd.DataFrame):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    target_excel = EXCEL_OUTPUT
    print(f"Generating enhanced Excel workbook at {target_excel}...")
    
    with pd.ExcelWriter(target_excel, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Master Knowledge Mine', index=False)
        
        golden_masters = df[df["Asset Format"] == "Golden Master Playbook"]
        golden_masters.to_excel(writer, sheet_name='🏆 Golden Master Playbooks', index=False)
        
        half_baked = df[df["Actionable Rating"].str.contains("Half-Baked|R&D", case=False, na=False)]
        half_baked.to_excel(writer, sheet_name='💡 Half-Baked Ideas', index=False)
        
        prompts = df[df["Primary Topic"].str.contains("Prompt", case=False, na=False) | df["Asset Title"].str.contains("Prompt", case=False, na=False)]
        prompts.to_excel(writer, sheet_name='🤖 Prompts & Engineering', index=False)
        
        automations = df[df["Asset Format"].str.contains("JSON|Workflow", case=False, na=False) | df["Primary Topic"].str.contains("Automation|Agent", case=False, na=False)]
        automations.to_excel(writer, sheet_name='⚡ Workable Automations', index=False)
        
        trading = df[df["Primary Topic"].str.contains("Trading|Finance", case=False, na=False)]
        trading.to_excel(writer, sheet_name='📈 Trading & Systems', index=False)
        
        vault = df[df["Actionable Rating"].str.contains("Sensitive|Private", case=False, na=False)]
        vault.to_excel(writer, sheet_name='🔒 Private Vault', index=False)
        
        kpi_topic = df["Primary Topic"].value_counts().reset_index()
        kpi_topic.columns = ["Primary Topic / Domain", "Asset Count"]
        kpi_topic.to_excel(writer, sheet_name='Executive KPI Summary', index=False, startrow=1, startcol=1)
        
        kpi_work = df["Workability Status"].value_counts().reset_index()
        kpi_work.columns = ["Workability Status", "Asset Count"]
        kpi_work.to_excel(writer, sheet_name='Executive KPI Summary', index=False, startrow=1, startcol=4)
        
        kpi_rating = df["Actionable Rating"].value_counts().reset_index()
        kpi_rating.columns = ["Actionable Rating", "Asset Count"]
        kpi_rating.to_excel(writer, sheet_name='Executive KPI Summary', index=False, startrow=1, startcol=7)

    wb = openpyxl.load_workbook(target_excel)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        if sheet_name == 'Executive KPI Summary':
            ws.cell(row=1, column=2, value="EXECUTIVE METRICS: SECOND BRAIN KNOWLEDGE MINE").font = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
        else:
            style_excel_sheet(ws)
            ws.auto_filter.ref = ws.dimensions
    wb.save(target_excel)
    print(f"Master Second Brain Excel workbook saved successfully to {target_excel}!")

    try:
        wb.save(EXCEL_OUTPUT_ALT)
        print(f"Also synced to {EXCEL_OUTPUT_ALT}")
    except Exception:
        pass

if __name__ == "__main__":
    df_master = process_all_files()
    build_enhanced_excel(df_master)
    print("\nALL DASHBOARD AND INDEX PIPELINES EXECUTED SUCCESSFULLY!")
